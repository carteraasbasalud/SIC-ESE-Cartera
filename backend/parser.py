import openpyxl
import sqlite3
import pandas as pd
import datetime
import os
import sys

# Add backend directory to path
sys.path.append(os.path.dirname(os.path.abspath(__file__)))
from db import get_db_connection, init_db, DB_PATH

EXCEL_PATH = r"g:\PROYECTOS-IA\Cartera\ESTADO CARTERA CORTE DIC 2025.xlsx"

def get_category_and_color(sheet_name, cell, obs, nit, entidad):
    if sheet_name == 'CONSOLIDADO CORTE DIC 2025':
        return "CONSOLIDADO", None
    elif sheet_name == 'ASMETSALUD YA REPORTADAS':
        return "ASMETSALUD YA REPORTADAS", None
    elif sheet_name == ' OTRAS EPS CORTE 2013-2022':
        return "OTRAS EPS CORTE 2013-2022", None
        
    # Analyze OTRAS EPS 2023
    color_hex = None
    fill = cell.fill
    fg_color = fill.fgColor
    if fg_color:
        if fg_color.type == 'rgb':
            color_hex = fg_color.rgb
        elif fg_color.type == 'indexed':
            color_hex = f"indexed_{fg_color.indexed}"
        elif fg_color.type == 'theme':
            color_hex = f"theme_{fg_color.theme}_tint_{fg_color.tint}"
            
    obs_upper = str(obs).upper() if obs else ""
    entidad_upper = str(entidad).upper() if entidad else ""
    
    # Check observations and color to map to the 14 categories from RTF
    if "CUOTA" in obs_upper or "PENSIONAL" in obs_upper or "CUOTAS" in obs_upper:
        category = "CUOTAS PARTES"
    elif "LIQUIDADA" in obs_upper or "LIQUIDACION" in obs_upper or "LIQ" in obs_upper or "CONCURSAL" in obs_upper:
        category = "EPS LIQUIDADAS"
    elif "INCAPACIDAD" in obs_upper or "INCAPACIDADES" in obs_upper or color_hex == "theme_2_tint_-0.249977111117893":
        category = "INCAPACIDADES"
    elif "VACUNACION" in obs_upper or "PAI" in obs_upper or "VIRAL" in obs_upper or color_hex == "theme_3_tint_0.3999755851924192":
        category = "VACUNACIÓN"
    elif "RETENCION" in obs_upper:
        category = "SALDO POR RETENCIÓN"
    elif "CAPITA" in obs_upper or "CAPITADO" in obs_upper or color_hex == "theme_5_tint_-0.249977111117893":
        category = "SALDO LIQUIDACIÓN CONTRACTUAL CÁPITA"
    elif "DEVOLUCION" in obs_upper or "DEVOLUCIONES" in obs_upper or color_hex == "theme_0_tint_-0.3499862666707358":
        category = "DEVOLUCIONES PARA PROCESO DE CONCILIACIÓN"
    elif ("REPORTE" in obs_upper and "PAGO" in obs_upper) or "SIN APLICACION" in obs_upper or color_hex == "FF99CCFF":
        category = "REPORTES DE PAGO, SIN APLICACIÓN EN CARTERA"
    elif "CONCILIACION GLOSA" in obs_upper or "GLOSAS" in obs_upper or color_hex == "FFFFFF00":
        category = "FACTURAS PARA CONCILIACIÓN (GESTIÓN), GLOSAS. CARTERA"
    elif "PRESCRIPCION" in obs_upper or "CORTE AL 31/07/2023" in obs_upper or color_hex == "FF92D050":
        category = "FACTURAS PARA INICIAR PROCESO DE TRÁMITE O PRESCRIPCIÓN ANTE LA EPS"
    elif "AGOSTO A DICIEMBRE DE 2023" in obs_upper or (color_hex == "00000000" or color_hex is None):
        category = "FACTURAS PARA INICIO DE TRÁMITE ANTE LA EPS"
    elif color_hex == "FFFA9986":
        category = "FACTURAS CONCILIADAS"
    else:
        category = "FACTURAS PARA INICIO DE TRÁMITE ANTE LA EPS" # default fallback
        
    return category, color_hex

def clean_value(val):
    if pd.isna(val):
        return None
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.isoformat()
    return val

def import_excel_to_db():
    print("Initializing Database...")
    init_db()
    
    conn = get_db_connection()
    cursor = conn.cursor()
    
    # Get existing invoices to check for updates
    print("Loading existing invoices from DB...")
    cursor.execute("SELECT id, consecutiv, nit, sheet_name, categoria, color_hex, observaciones FROM invoices;")
    existing_rows = [dict(r) for r in cursor.fetchall()]
    
    # Map (consecutiv, nit, sheet_name) -> db_id
    existing_map = {}
    for r in existing_rows:
        key = (r['consecutiv'], r['nit'], r['sheet_name'])
        existing_map[key] = r['id']
        
    # Map (consecutiv, nit) -> list of existing rows
    existing_map_by_key = {}
    for r in existing_rows:
        k = (r['consecutiv'], r['nit'])
        if k not in existing_map_by_key:
            existing_map_by_key[k] = []
        existing_map_by_key[k].append(r)
        
    # Detect file type by opening with openpyxl
    print("Opening Excel with openpyxl to detect format...")
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True)
    sheet_names = wb.sheetnames
    is_report_81 = 'Hoja1' in sheet_names and len(sheet_names) == 1
    wb.close()
    
    processed_keys = set()
    invoices_to_insert = []
    invoices_to_update = []
    
    if is_report_81:
        print("Detected format: REPORTE 81 (Single-sheet unified platform report)")
        # Load Hoja1 with Pandas skipping the first 10 rows (header is at index 10)
        df = pd.read_excel(EXCEL_PATH, sheet_name='Hoja1', header=10)
        print(f"Loaded {df.shape[0]} rows from Hoja1.")
        
        for idx, row in df.iterrows():
            row_nit = row.get('NIT')
            row_entidad = row.get('CLIENTE')
            
            if pd.isna(row_nit) and pd.isna(row_entidad):
                continue
                
            nit = str(row_nit).split('.')[0] if not pd.isna(row_nit) else None
            entidad = str(row_entidad).strip() if not pd.isna(row_entidad) else None
            cta1 = str(row.get('CODIGO CONTABLE')).split('.')[0] if not pd.isna(row.get('CODIGO CONTABLE')) else None
            tipo = str(row.get('TIPO')).split('.')[0] if not pd.isna(row.get('TIPO')) else None
            consecutiv = str(row.get('NUMERO FACTURA')).split('.')[0] if not pd.isna(row.get('NUMERO FACTURA')) else None
            
            fvence = clean_value(row.get('VENCIMIENTO'))
            fecha_vencimiento = fvence
            fecha_emision = clean_value(row.get('FECHA FACTURA'))
            radicado = clean_value(row.get('RADICADO'))
            
            valor = row.get('VALOR INICIAL')
            abonos = row.get('ABONO')
            saldo = row.get('SALDO')
            
            valor = float(valor) if pd.notna(valor) else 0.0
            abonos = float(abonos) if pd.notna(abonos) else 0.0
            saldo = float(saldo) if pd.notna(saldo) else 0.0
            
            ano = None
            if fecha_emision:
                try:
                    if isinstance(fecha_emision, str):
                        ano = int(fecha_emision.split('-')[0])
                    elif isinstance(fecha_emision, (datetime.datetime, datetime.date)):
                        ano = fecha_emision.year
                except:
                    pass
            
            nodias = row.get('DIAS')
            nodias = int(nodias) if pd.notna(nodias) else None
            
            porven = float(row.get('POR VENCER')) if 'POR VENCER' in row and pd.notna(row.get('POR VENCER')) else 0.0
            ven060 = float(row.get('VENCIDO 60')) if 'VENCIDO 60' in row and pd.notna(row.get('VENCIDO 60')) else 0.0
            ven090 = float(row.get('VENCIDO 90')) if 'VENCIDO 90' in row and pd.notna(row.get('VENCIDO 90')) else 0.0
            ven180 = float(row.get('VENCIDO 180')) if 'VENCIDO 180' in row and pd.notna(row.get('VENCIDO 180')) else 0.0
            ven360 = float(row.get('VENCIDO 360')) if 'VENCIDO 360' in row and pd.notna(row.get('VENCIDO 360')) else 0.0
            venmas = float(row.get('VENCIDOMA DE 360')) if 'VENCIDOMA DE 360' in row and pd.notna(row.get('VENCIDOMA DE 360')) else 0.0
            
            # Find if this (consecutiv, nit) exists in the DB under any sheet name
            k = (consecutiv, nit)
            if k in existing_map_by_key:
                # Update all matching records in DB
                for r in existing_map_by_key[k]:
                    db_id = r['id']
                    s_id = r['sheet_name']
                    category = r['categoria']
                    color_hex = r['color_hex']
                    observaciones = r['observaciones']
                    
                    invoices_to_update.append((
                        entidad, cta1, tipo, fvence, valor, abonos, saldo,
                        ano, fecha_emision, radicado, nodias, porven, ven060, ven090,
                        ven180, ven360, venmas, category, color_hex, observaciones, db_id
                    ))
                    processed_keys.add((consecutiv, nit, s_id))
            else:
                # Insert as a new invoice under CONSOLIDADO
                s_id = 'CONSOLIDADO'
                category = 'CONSOLIDADO'
                color_hex = None
                observaciones = None
                
                invoices_to_insert.append((
                    s_id, nit, entidad, cta1, tipo, consecutiv, fvence,
                    valor, abonos, saldo, ano, fecha_emision, radicado, nodias,
                    porven, ven060, ven090, ven180, ven360, venmas,
                    category, color_hex, observaciones
                ))
                processed_keys.add((consecutiv, nit, s_id))
                
        # Batch insert and update
        if invoices_to_insert:
            print(f"Inserting {len(invoices_to_insert)} new invoices...")
            cursor.executemany("""
            INSERT INTO invoices (
                sheet_name, nit, nombre, cta1, tipo, consecutiv, fecha_vencimiento,
                valor, abonos, saldo, ano, fecha_emision, radicado, nodias,
                porven, ven060, ven090, ven180, ven360, venmas,
                categoria, color_hex, observaciones
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
            """, invoices_to_insert)
            conn.commit()
            
        if invoices_to_update:
            print(f"Updating {len(invoices_to_update)} existing invoices...")
            cursor.executemany("""
            UPDATE invoices SET
                nombre = ?, cta1 = ?, tipo = ?, fecha_vencimiento = ?, valor = ?, abonos = ?, saldo = ?,
                ano = ?, fecha_emision = ?, radicado = ?, nodias = ?, porven = ?, ven060 = ?, ven090 = ?,
                ven180 = ?, ven360 = ?, venmas = ?, categoria = ?, color_hex = ?, observaciones = ?
            WHERE id = ?;
            """, invoices_to_update)
            conn.commit()
            
    else:
        print("Detected format: Initial Load Multi-Sheet (ESTADO CARTERA CORTE DIC 2025)")
        # Load with openpyxl to get cell colors for OTRAS EPS 2023
        print("Opening Excel with openpyxl to read sheet 'OTRAS EPS 2023' styles...")
        wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
        sheet_2023 = wb['OTRAS EPS 2023']
        row_colors_2023 = {}
        
        for r in range(2, sheet_2023.max_row + 1):
            cell_entidad = sheet_2023.cell(row=r, column=2) # ENTIDAD column
            obs = sheet_2023.cell(row=r, column=13).value
            nit = sheet_2023.cell(row=r, column=1).value
            entidad = cell_entidad.value
            
            if nit is None and entidad is None:
                continue
                
            category, color_hex = get_category_and_color('OTRAS EPS 2023', cell_entidad, obs, nit, entidad)
            row_colors_2023[r] = (category, color_hex)
            
        wb.close()
        print("Excel styles loaded. Now reading all sheets with Pandas...")
        
        xl = pd.ExcelFile(EXCEL_PATH)
        for sheet_name in xl.sheet_names:
            print(f"\nProcessing sheet: {sheet_name}...")
            df = xl.parse(sheet_name)
            print(f"Loaded {df.shape[0]} rows.")
            
            invoices_to_insert = []
            invoices_to_update = []
            
            for idx, row in df.iterrows():
                row_nit = row.get('nit') if 'nit' in row else row.get('NIT')
                row_entidad = row.get('nombre') if 'nombre' in row else row.get('ENTIDAD')
                
                if pd.isna(row_nit) and pd.isna(row_entidad):
                    continue
                    
                nit = str(row_nit).split('.')[0] if not pd.isna(row_nit) else None
                entidad = str(row_entidad).strip() if not pd.isna(row_entidad) else None
                cta1 = str(row.get('cta1')).split('.')[0] if not pd.isna(row.get('cta1')) else None
                tipo = str(row.get('tipo')).split('.')[0] if not pd.isna(row.get('tipo')) else None
                
                consecutiv = str(row.get('consecutiv')).split('.')[0] if not pd.isna(row.get('consecutiv')) else None
                if not consecutiv and 'FACTURA HOMOLOGADA' in row:
                    consecutiv = str(row.get('FACTURA HOMOLOGADA')).split('.')[0] if not pd.isna(row.get('FACTURA HOMOLOGADA')) else None
                    
                fvence = clean_value(row.get('fvence') if 'fvence' in row else row.get('FECHA VENCIMIENTO'))
                fecha_vencimiento = fvence
                
                fecha_emision = clean_value(row.get('fecha'))
                radicado = clean_value(row.get('radicado'))
                
                valor = row.get('valor') if 'valor' in row else row.get('VALOR FACTURA')
                abonos = row.get('abonos') if 'abonos' in row else row.get('ABONOS')
                saldo = row.get('saldo') if 'saldo' in row else row.get('SALDO PENDIENTE')
                
                valor = float(valor) if pd.notna(valor) else 0.0
                abonos = float(abonos) if pd.notna(abonos) else 0.0
                saldo = float(saldo) if pd.notna(saldo) else 0.0
                
                ano = row.get('AO') if 'AO' in row else row.get('AÑO')
                if pd.isna(ano):
                    ano = None
                else:
                    try:
                        ano = int(float(ano))
                    except:
                        ano = None
                        
                nodias = row.get('nodias') if 'nodias' in row else None
                nodias = int(nodias) if pd.notna(nodias) else None
                
                porven = float(row.get('porven')) if 'porven' in row and pd.notna(row.get('porven')) else 0.0
                ven060 = float(row.get('ven060')) if 'ven060' in row and pd.notna(row.get('ven060')) else 0.0
                ven090 = float(row.get('ven090')) if 'ven090' in row and pd.notna(row.get('ven090')) else 0.0
                ven180 = float(row.get('ven180')) if 'ven180' in row and pd.notna(row.get('ven180')) else 0.0
                ven360 = float(row.get('ven360')) if 'ven360' in row and pd.notna(row.get('ven360')) else 0.0
                venmas = float(row.get('venmas')) if 'venmas' in row and pd.notna(row.get('venmas')) else 0.0
                
                observaciones = str(row.get('OBSERVACIONES')) if 'OBSERVACIONES' in row and pd.notna(row.get('OBSERVACIONES')) else None
                
                if sheet_name == 'CONSOLIDADO CORTE DIC 2025':
                    s_id = 'CONSOLIDADO'
                    category = 'CONSOLIDADO'
                    color_hex = None
                elif sheet_name == 'ASMETSALUD YA REPORTADAS':
                    s_id = 'ASMETSALUD_REPORTADAS'
                    category = 'ASMETSALUD REPORTADAS'
                    color_hex = None
                elif sheet_name == ' OTRAS EPS CORTE 2013-2022':
                    s_id = 'OTRAS_EPS_2013_2022'
                    category = 'OTRAS EPS CORTE 2013-2022'
                    color_hex = None
                elif sheet_name == 'OTRAS EPS 2023':
                    s_id = 'OTRAS_EPS_2023'
                    excel_row_idx = idx + 2
                    category, color_hex = row_colors_2023.get(excel_row_idx, ('OTRAS FACTURAS', None))
                else:
                    s_id = sheet_name
                    category = 'OTRAS'
                    color_hex = None
                    
                key = (consecutiv, nit, s_id)
                processed_keys.add(key)
                
                # Check if exists in DB
                if key in existing_map:
                    db_id = existing_map[key]
                    invoices_to_update.append((
                        entidad, cta1, tipo, fvence, valor, abonos, saldo,
                        ano, fecha_emision, radicado, nodias, porven, ven060, ven090,
                        ven180, ven360, venmas, category, color_hex, observaciones, db_id
                    ))
                else:
                    invoices_to_insert.append((
                        s_id, nit, entidad, cta1, tipo, consecutiv, fvence,
                        valor, abonos, saldo, ano, fecha_emision, radicado, nodias,
                        porven, ven060, ven090, ven180, ven360, venmas,
                        category, color_hex, observaciones
                    ))
                    
            # Batch insert and update
            if invoices_to_insert:
                print(f"Inserting {len(invoices_to_insert)} new invoices...")
                cursor.executemany("""
                INSERT INTO invoices (
                    sheet_name, nit, nombre, cta1, tipo, consecutiv, fecha_vencimiento,
                    valor, abonos, saldo, ano, fecha_emision, radicado, nodias,
                    porven, ven060, ven090, ven180, ven360, venmas,
                    categoria, color_hex, observaciones
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
                """, invoices_to_insert)
                conn.commit()
                
            if invoices_to_update:
                print(f"Updating {len(invoices_to_update)} existing invoices...")
                cursor.executemany("""
                UPDATE invoices SET
                    nombre = ?, cta1 = ?, tipo = ?, fecha_vencimiento = ?, valor = ?, abonos = ?, saldo = ?,
                    ano = ?, fecha_emision = ?, radicado = ?, nodias = ?, porven = ?, ven060 = ?, ven090 = ?,
                    ven180 = ?, ven360 = ?, venmas = ?, categoria = ?, color_hex = ?, observaciones = ?
                WHERE id = ?;
                """, invoices_to_update)
                conn.commit()
                
    # Set missing ones to saldo = 0 (paid)
    unprocessed_keys = set(existing_map.keys()) - processed_keys
    if unprocessed_keys:
        print(f"\nProcessing {len(unprocessed_keys)} paid/missing invoices...")
        ids_to_zero = [existing_map[k] for k in unprocessed_keys]
        # SQLite limit for variables is 999, so we do it in batches
        batch_size = 500
        for i in range(0, len(ids_to_zero), batch_size):
            batch = ids_to_zero[i:i+batch_size]
            placeholders = ",".join(["?"] * len(batch))
            cursor.execute(f"""
                UPDATE invoices 
                SET saldo = 0.0, observaciones = 'Saldada (no figura en el último reporte cargado)'
                WHERE id IN ({placeholders});
            """, batch)
            conn.commit()
        print("Updated missing invoices to 0 balance.")

    conn.close()
    print("\nExcel parsing and smart database sync finished successfully!")

if __name__ == "__main__":
    import_excel_to_db()
