import openpyxl

wb = openpyxl.load_workbook(r'G:\PROYECTOS-IA\Cartera\Dcts\COMPARADO JULIO 2026-2025-1.xlsx', data_only=True)
ws = wb['Hoja1']

data = []
for r in range(6, 20):
    concepto = ws.cell(row=r, column=3).value
    val_2026 = ws.cell(row=r, column=4).value or 0
    part_2026 = ws.cell(row=r, column=5).value or 0
    val_2025 = ws.cell(row=r, column=7).value or 0
    part_2025 = ws.cell(row=r, column=8).value or 0
    var_pct = ws.cell(row=r, column=9).value or 0
    
    val_2026 = float(val_2026) if val_2026 != '' else 0.0
    val_2025 = float(val_2025) if val_2025 != '' else 0.0
    diff_abs = val_2026 - val_2025
    real_pct = (diff_abs / val_2025 * 100) if val_2025 != 0 else (100.0 if val_2026 > 0 else 0.0)
    
    data.append({
        'concepto': concepto.strip() if concepto else 'SIN CONCEPTO',
        'val_2026': val_2026,
        'part_2026': float(part_2026) if part_2026 != '' else 0.0,
        'val_2025': val_2025,
        'part_2025': float(part_2025) if part_2025 != '' else 0.0,
        'diff_abs': diff_abs,
        'var_pct_excel': float(var_pct) if var_pct != '' else 0.0,
        'real_pct': real_pct
    })

print(f"{'CONCEPTO':<35} | {'JULIO 2025':>18} | {'PART 25':>8} | {'JULIO 2026':>18} | {'PART 26':>8} | {'VAR ABSOLUTA':>18} | {'VAR %':>8}")
print('-' * 125)
for d in data:
    c = d['concepto']
    v25 = f"${d['val_2025']:,.2f}"
    p25 = f"{d['part_2025']:.2f}%"
    v26 = f"${d['val_2026']:,.2f}"
    p26 = f"{d['part_2026']:.2f}%"
    vdiff = f"${d['diff_abs']:,.2f}"
    pdiff = f"{d['real_pct']:+.2f}%"
    print(f"{c:<35} | {v25:>18} | {p25:>8} | {v26:>18} | {p26:>8} | {vdiff:>18} | {pdiff:>8}")

tot_2025 = sum(d['val_2025'] for d in data)
tot_2026 = sum(d['val_2026'] for d in data)
diff_tot = tot_2026 - tot_2025
pct_tot = (diff_tot / tot_2025) * 100
print('-' * 125)
print(f"{'TOTAL CARTERA NETA':<35} | {f'${tot_2025:,.2f}':>18} | {'100.00%':>8} | {f'${tot_2026:,.2f}':>18} | {'100.00%':>8} | {f'${diff_tot:,.2f}':>18} | {f'{pct_tot:+.2f}%':>8}")
